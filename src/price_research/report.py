from __future__ import annotations

from datetime import datetime
import os
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .models import Product, StoreResult
from .price_evidence import eligible_price

HEADERS = ["DESCRIÇÃO DO PRODUTO", "COOCERQUI", "PAGUE MENOS"]


def _value(result, key: str, default=None):
    if isinstance(result, dict):
        return result.get(key, default)
    return getattr(result, key, default)


def report_price(result: StoreResult | dict) -> float | str:
    """Expose the cheapest valid displayed price after evidence validation.

    Club/CPF/Pix/coupon prices remain eligible by design. The operational sheet
    therefore continues to show the cheapest valid offer rather than forcing a
    distinction between normal and loyalty prices.
    """
    return eligible_price(result)


def reference_price(product: Product, result: StoreResult) -> float | None:
    """Return only a price already expressed in the requested sale unit."""
    price = report_price(result)
    if not isinstance(price, (int, float)):
        return None
    result_unit = (_value(result, "unit_price", "") or product.unidade_preco).lower()
    if result_unit != product.unidade_preco.lower():
        return None
    return float(price)


def save_report(
    records,
    output_dir: Path,
    timezone: str = "America/Sao_Paulo",
    *,
    complete: bool = True,
    total_products: int | None = None,
) -> Path:
    """Generate the three-column operational spreadsheet.

    A partial run is unmistakable from its filename; the workbook layout stays
    exactly as requested by the operator.
    """
    output_dir.mkdir(parents=True, exist_ok=True)
    now = datetime.now(ZoneInfo(timezone))
    stamp = now.strftime("%Y-%m-%d_%H-%M-%S")
    if complete:
        filename = f"pesquisa_precos_{stamp}.xlsx"
    else:
        total = total_products if isinstance(total_products, int) and total_products > 0 else "?"
        filename = f"pesquisa_precos_{stamp}_PARCIAL_{len(records)}-de-{total}.xlsx"
    path = output_dir / filename

    wb = Workbook()
    ws = wb.active
    ws.title = "Pesquisa de Preços"
    ws.sheet_view.showGridLines = False
    ws.append(HEADERS)
    for product, coo, pag in records:
        ws.append([product.search_name, report_price(coo), report_price(pag)])

    header_fill = PatternFill("solid", fgColor="30343B")
    white_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D1D5DB")
    section = Side(style="medium", color="6B7280")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=section)
    ws.row_dimensions[1].height = 26

    previous_category = None
    for row_index, (product, _coo, _pag) in enumerate(records, start=2):
        category_changed = previous_category is not None and product.categoria != previous_category
        for column in range(1, 4):
            cell = ws.cell(row_index, column)
            cell.border = Border(top=section if category_changed else None, bottom=thin)
            cell.alignment = Alignment(
                horizontal="left" if column == 1 else "right",
                vertical="center",
            )
            if column in (2, 3):
                if isinstance(cell.value, (int, float)):
                    cell.number_format = 'R$ #,##0.00'
                else:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row_index].height = 21
        previous_category = product.categoria

    ws.column_dimensions["A"].width = 62
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    temporary = path.with_name(path.stem + ".tmp.xlsx")
    try:
        wb.save(temporary)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)
    return path
