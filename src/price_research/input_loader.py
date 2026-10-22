from __future__ import annotations

import csv
import json
import unicodedata
from collections import defaultdict, deque
from pathlib import Path
from openpyxl import load_workbook

from .models import Product

ALIASES = {
    "codigo": "codigo_interno",
    "código": "codigo_interno",
    "codigo interno": "codigo_interno",
    "código interno": "codigo_interno",
    "produto": "descricao_original",
    "descricao do produto": "descricao_original",
    "descrição do produto": "descricao_original",
    "descricao": "descricao_original",
    "descrição": "descricao_original",
    "nome": "nome_normalizado",
}


def _norm_header(value) -> str:
    return " ".join(str(value or "").strip().lower().split())


def _split_terms(value) -> list[str]:
    if value is None:
        return []
    if isinstance(value, (list, tuple, set)):
        return [str(x).strip() for x in value if str(x).strip()]
    return [x.strip() for x in str(value).split(";") if x.strip()]


def _match_key(value) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return " ".join(text.casefold().split())


def _row_to_product(row: dict) -> Product:
    normalized = {}
    for key, value in row.items():
        k = _norm_header(key)
        k = ALIASES.get(k, k.replace(" ", "_"))
        normalized[k] = value
    return Product(
        codigo_interno=str(normalized.get("codigo_interno") or "").strip(),
        descricao_original=str(normalized.get("descricao_original") or "").strip(),
        nome_normalizado=str(normalized.get("nome_normalizado") or "").strip(),
        marca=str(normalized.get("marca") or "").strip(),
        categoria=str(normalized.get("categoria") or "").strip(),
        termos_obrigatorios=_split_terms(normalized.get("termos_obrigatorios")),
        termos_substituiveis=_split_terms(normalized.get("termos_substituiveis")),
        termos_excluidos=_split_terms(normalized.get("termos_excluidos")),
        aliases=_split_terms(normalized.get("aliases")),
        aliases_coocerqui=_split_terms(normalized.get("aliases_coocerqui")),
        aliases_paguemenos=_split_terms(normalized.get("aliases_paguemenos")),
        unidade_preco=str(normalized.get("unidade_preco") or "un").strip(),
        observacoes=str(normalized.get("observacoes") or "").strip(),
    )


def _load_companion_rules(path: Path) -> dict[str, deque[Product]]:
    """Lê regras técnicas sem expô-las na planilha usada pelo operador."""
    rules_path = path.with_name(f"{path.stem}_regras.json")
    if not rules_path.exists():
        return {}
    try:
        payload = json.loads(rules_path.read_text(encoding="utf-8"))
        raw_products = payload.get("products", []) if isinstance(payload, dict) else []
        queues: dict[str, deque[Product]] = defaultdict(deque)
        for raw in raw_products:
            if not isinstance(raw, dict):
                continue
            display_name = raw.get("produto_planilha") or raw.get("nome_normalizado") or raw.get("descricao_original")
            product = _row_to_product(raw)
            if display_name and product.search_name:
                queues[_match_key(display_name)].append(product)
        return queues
    except (OSError, json.JSONDecodeError, TypeError, ValueError) as exc:
        raise ValueError(f"Arquivo de regras inválido: {rules_path}") from exc


def _load_alias_overrides(path: Path) -> dict[str, dict[str, list[str]]]:
    """Lê equivalências de nomenclatura revisadas separadamente por loja."""
    aliases_path = path.with_name(f"{path.stem}_aliases.json")
    if not aliases_path.exists():
        return {}
    try:
        payload = json.loads(aliases_path.read_text(encoding="utf-8"))
        raw_products = payload.get("products", []) if isinstance(payload, dict) else []
        overrides: dict[str, dict[str, list[str]]] = {}
        for raw in raw_products:
            if not isinstance(raw, dict):
                continue
            display_name = raw.get("produto_planilha") or raw.get("nome_normalizado")
            if not display_name:
                continue
            overrides[_match_key(display_name)] = {
                "aliases": _split_terms(raw.get("aliases")),
                "aliases_coocerqui": _split_terms(raw.get("aliases_coocerqui")),
                "aliases_paguemenos": _split_terms(raw.get("aliases_paguemenos")),
            }
        return overrides
    except (OSError, json.JSONDecodeError, TypeError, ValueError) as exc:
        raise ValueError(f"Arquivo de aliases inválido: {aliases_path}") from exc


def _merge_unique(current: list[str], extra: list[str]) -> list[str]:
    merged = list(current)
    for value in extra:
        if value and value not in merged:
            merged.append(value)
    return merged


def _enrich_from_companion(products: list[Product], path: Path) -> list[Product]:
    rules = _load_companion_rules(path)
    aliases = _load_alias_overrides(path)
    if not rules and not aliases:
        return products
    enriched: list[Product] = []
    for product in products:
        candidates = rules.get(_match_key(product.search_name))
        selected = candidates.popleft() if candidates else product
        override = aliases.get(_match_key(selected.search_name), {})
        selected.aliases = _merge_unique(selected.aliases, override.get("aliases", []))
        selected.aliases_coocerqui = _merge_unique(
            selected.aliases_coocerqui, override.get("aliases_coocerqui", [])
        )
        selected.aliases_paguemenos = _merge_unique(
            selected.aliases_paguemenos, override.get("aliases_paguemenos", [])
        )
        enriched.append(selected)
    return enriched


def load_products(path: str | Path) -> list[Product]:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)
    rows: list[dict] = []
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            rows = list(csv.DictReader(f))
    elif path.suffix.lower() in {".xlsx", ".xlsm"}:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb["Catalogo de Entrada"] if "Catalogo de Entrada" in wb.sheetnames else wb.active
        values = list(ws.iter_rows(values_only=True))
        header_idx = None
        for i, row in enumerate(values):
            headers = [_norm_header(v) for v in row]
            if "codigo_interno" in headers or "codigo interno" in headers or "descricao_original" in headers:
                header_idx = i
                break
        if header_idx is None:
            header_idx = 0
        headers = [str(v or "").strip() for v in values[header_idx]]
        for row in values[header_idx + 1:]:
            if not any(v not in (None, "") for v in row):
                continue
            rows.append(dict(zip(headers, row)))
    else:
        raise ValueError("Formato de entrada suportado: .xlsx ou .csv")

    products = [_row_to_product(r) for r in rows]
    products = [p for p in products if p.search_name]
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        products = _enrich_from_companion(products, path)
    if not products:
        raise ValueError("Nenhum produto válido encontrado no arquivo de entrada")
    return products
